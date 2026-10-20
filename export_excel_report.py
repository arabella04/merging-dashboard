import os
import json
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

def export_styled_excel(output_filename="QA_Audit_Trail_Report.xlsx", risk_filter=None, operators=None):
    json_path = os.path.join(BASE_DIR, "data.json")
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    records = data.get("records", [])
    if risk_filter:
        records = [r for r in records if r["risk_level"] in risk_filter]
    if operators:
        records = [r for r in records if r["qa_operator"] in operators]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "QA Audit Trail"

    headers = [
        "Audit ID",
        "Date/Time",
        "QA User",
        "Action",
        "Primary Account",
        "Affected Account",
        "Risk Level",
        "Discrepancies",
        "Status",
        "Auditor Notes"
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1E222D", end_color="1E222D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB")
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")

    for r_idx, rec in enumerate(records, start=2):
        row_vals = [
            rec["audit_id"],
            rec["timestamp_display"],
            rec["qa_operator"],
            rec["action"],
            f"{rec['primary_account']['formatted']} (DOB: {rec['primary_account']['dob']})",
            f"{rec['affected_account']['formatted']} (DOB: {rec['affected_account']['dob']})",
            rec["risk_level"],
            rec["discrepancy_text"],
            rec["status"],
            rec.get("auditor_notes", "")
        ]
        ws.append(row_vals)

        risk = rec["risk_level"]
        if risk == "HIGH RISK":
            row_fill = PatternFill(start_color="FDE8E8", end_color="FDE8E8", fill_type="solid")
            badge_fill = PatternFill(start_color="DC3545", end_color="DC3545", fill_type="solid")
        elif risk == "MEDIUM / REVIEW":
            row_fill = PatternFill(start_color="FFF8E6", end_color="FFF8E6", fill_type="solid")
            badge_fill = PatternFill(start_color="FD7E14", end_color="FD7E14", fill_type="solid")
        else:
            row_fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
            badge_fill = PatternFill(start_color="28A745", end_color="28A745", fill_type="solid")

        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=(c_idx in [5, 6, 8, 10]))
            if c_idx == 7:
                cell.fill = badge_fill
                cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.fill = row_fill
                cell.font = Font(name="Calibri", size=10, color="0F172A")

    # Auto-fit columns
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 58)

    ws.freeze_panes = "A2"
    out_path = os.path.join(BASE_DIR, output_filename)
    wb.save(out_path)
    print(f"Saved styled Excel report to {out_path} ({len(records)} rows)")
    return out_path

if __name__ == "__main__":
    export_styled_excel("QA_Audit_Trail_Full_Report.xlsx")
    export_styled_excel("QA_Audit_Trail_High_Risk_Only.xlsx", risk_filter=["HIGH RISK"])
